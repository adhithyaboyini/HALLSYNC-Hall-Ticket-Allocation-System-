"""
step2_create_school_template.py  —  HallSync Step 2
Creates a clean manual-entry School_Info.xlsx template for admin.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# ── Colours ───────────────────────────────────────────────────────────────────
NAVY   = "1A3C6E";  BLUE  = "2E6DA4";  GREEN = "0D5C2F"
LIGHT  = "EEF2F9";  WHITE = "FFFFFF";  GOLD  = "FFC107"
GREY   = "555555";  YBGD  = "FFF8E7";  GBGD  = "E8F5E9"
YELLOW = "FFFDE7"

def tb(c="BBBBBB"):
    s = Side(style="thin", color=c)
    return Border(left=s, right=s, top=s, bottom=s)

def hc(cell, text, bg=NAVY, fg=WHITE, sz=10, wrap=False, bold=True):
    cell.value     = text
    cell.font      = Font(bold=bold, size=sz, color=fg, name="Arial")
    cell.fill      = PatternFill("solid", start_color=bg)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=wrap)
    cell.border    = tb()

def dc(cell, bg=WHITE, bold=False, color="000000", center=True, sz=9):
    cell.fill      = PatternFill("solid", start_color=bg)
    cell.alignment = Alignment(horizontal="center" if center else "left",
                               vertical="center")
    cell.border    = tb()
    cell.font      = Font(name="Arial", size=sz, bold=bold, color=color)


wb = Workbook()

# ════════════════════════════════════════════════════════════════════
# SHEET 1 — School_Info
# ════════════════════════════════════════════════════════════════════
ws = wb.active
ws.title = "School_Info"

# Row 1 — Title
ws.merge_cells("A1:F1")
hc(ws["A1"], "🏫  HallSync — School Information Sheet", bg=NAVY, sz=14)
ws.row_dimensions[1].height = 32

# Row 2 — Subtitle
ws.merge_cells("A2:F2")
c2 = ws["A2"]
c2.value     = "Fill one row per school  •  All fields required  •  Do NOT rename or delete headers"
c2.font      = Font(italic=True, size=9, color=GREY, name="Arial")
c2.fill      = PatternFill("solid", start_color=YBGD)
c2.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[2].height = 16

# Row 3 — Gold accent
for col in range(1, 7):
    ws.cell(3, col).fill = PatternFill("solid", start_color=GOLD)
ws.row_dimensions[3].height = 4

# Row 4 — Section labels
ws.merge_cells("A4:D4")
hc(ws["A4"], "📋  School Details", bg=NAVY, sz=10)
ws.merge_cells("E4:F4")
hc(ws["E4"], "🎯  Centre Info", bg=GREEN, sz=10)
ws.row_dimensions[4].height = 20

# Row 5 — Column headers
COLS = [
    ("School_Code",   16, NAVY),
    ("School_Name",   35, NAVY),
    ("Area",          22, NAVY),
    ("District_Code", 18, NAVY),
    ("Can_Be_Centre", 16, GREEN),
    ("Capacity",      14, GREEN),
]
for i, (name, width, bg) in enumerate(COLS, 1):
    hc(ws.cell(5, i), name, bg=bg, sz=10, wrap=True)
    ws.column_dimensions[get_column_letter(i)].width = width
ws.row_dimensions[5].height = 28

# Rows 6–105 — 100 empty data rows
DATA_START, DATA_END = 6, 105
for row in range(DATA_START, DATA_END + 1):
    bg = LIGHT if row % 2 == 0 else WHITE
    # School detail cols (A–D) — plain entry
    for col in range(1, 5):
        dc(ws.cell(row, col), bg=bg)
    # Can_Be_Centre (E) — yellow, needs filling
    ce = ws.cell(row, 5)
    dc(ce, bg=YELLOW, bold=True, color=GREEN)
    # Capacity (F) — yellow, needs filling
    cf = ws.cell(row, 6)
    dc(cf, bg=YELLOW, bold=True, color=GREEN)
    ws.row_dimensions[row].height = 17

# Can_Be_Centre — strict Yes/No dropdown
dv = DataValidation(
    type="list", formula1='"Yes,No"',
    allow_blank=False, showDropDown=False,
    showErrorMessage=True,
    errorTitle="Invalid Value",
    error="Please select Yes or No from the dropdown."
)
ws.add_data_validation(dv)
dv.sqref = f"E{DATA_START}:E{DATA_END}"

# Summary row at bottom
sum_row = DATA_END + 2
ws.merge_cells(f"A{sum_row}:D{sum_row}")
sl = ws.cell(sum_row, 1, "Totals")
sl.font      = Font(bold=True, size=10, color=WHITE, name="Arial")
sl.fill      = PatternFill("solid", start_color=NAVY)
sl.alignment = Alignment(horizontal="center", vertical="center")
sl.border    = tb()

# Count of centres
cc = ws.cell(sum_row, 5)
cc.value     = f'=COUNTIF(E{DATA_START}:E{DATA_END},"Yes")&" centre(s) selected"'
cc.font      = Font(bold=True, size=9, color=WHITE, name="Arial")
cc.fill      = PatternFill("solid", start_color=GREEN)
cc.alignment = Alignment(horizontal="center", vertical="center")
cc.border    = tb()
ws.merge_cells(f"E{sum_row}:F{sum_row}")
ws.row_dimensions[sum_row].height = 22

ws.freeze_panes = "A6"

# ════════════════════════════════════════════════════════════════════
# SHEET 2 — Instructions
# ════════════════════════════════════════════════════════════════════
wi = wb.create_sheet("Instructions")
wi.sheet_view.showGridLines = False
wi.column_dimensions["A"].width = 4
wi.column_dimensions["B"].width = 22
wi.column_dimensions["C"].width = 62

wi.merge_cells("B1:C1")
hc(wi["B1"], "📋  School_Info — How to Fill", bg=NAVY, sz=13)
wi.row_dimensions[1].height = 30

fields = [
    ("FIELD",           "DESCRIPTION",                                                         True),
    ("School_Code",     "Unique code for the school.\nExample: SCH101, SCH102\nMust be unique for every row — no duplicates allowed.", False),
    ("School_Name",     "Full official name of the school.\nExample: Alpha Higher Secondary School", False),
    ("Area",            "Area or locality where the school is located.\nExample: Chennai North, Coimbatore East", False),
    ("District_Code",   "District code for the school.\nExample: TN-CHN-01, TN-CBE-02",       False),
    ("Can_Be_Centre",   "Dropdown → Yes or No\nYes  = This school can host exams for other school students.\nNo   = Students from this school will be sent to another centre.", False),
    ("Capacity",        "Maximum number of students this school can seat during exams.\nEnter 0 or leave blank if Can_Be_Centre = No.\nExample: 120", False),
    ("",                "",                                                                    False),
    ("⚠  Rules",       "• One row = one school\n• School_Code must be unique — no duplicates\n• At least one school must have Can_Be_Centre = Yes\n• Capacity must be a number — no text\n• Do NOT rename or delete column headers\n• Do NOT merge cells in the data area\n• Save as School_Info.xlsx before running Step 3", False),
]

for r, (f, d, ih) in enumerate(fields, 3):
    b = wi.cell(r, 2, f)
    c = wi.cell(r, 3, d)
    if ih:
        hc(b, f, bg=BLUE); hc(c, d, bg=BLUE)
    else:
        b.font = Font(bold=True, size=9, name="Arial")
        b.fill = PatternFill("solid", start_color=LIGHT)
        b.alignment = Alignment(vertical="top"); b.border = tb()
        c.font = Font(size=9, name="Arial")
        c.fill = PatternFill("solid", start_color=WHITE)
        c.alignment = Alignment(vertical="top", wrap_text=True); c.border = tb()
    wi.row_dimensions[r].height = 50 if "\n" in d else 20

# ════════════════════════════════════════════════════════════════════
# SAVE
# ════════════════════════════════════════════════════════════════════
OUT = "/home/claude/hallsync_v2/data/School_Info_Template.xlsx"
wb.save(OUT)
print(f"School_Info template saved → {OUT}")
